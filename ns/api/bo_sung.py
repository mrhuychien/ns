"""§4.7 — Bổ sung thông tin còn thiếu: tổng quan, lưới NV × field, lưu 1 ô,
mẫu Excel, xem trước & áp dụng Excel."""
import io
import re

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, today

from ns.api import common, guard

# Bộ field "hồ sơ tối thiểu" — thứ tự cố định = cột lưới = cột Excel.
# ap_dung: 'tat_ca' | 'tru_khoan' | 'dong_bhxh' | 'khong_bhxh' | 'co_luong_tk'
TRUONG_TOI_THIEU = [
    {"nhom": "Cá nhân", "field": "custom_cccd", "ap_dung": "tat_ca"},
    {"nhom": "Cá nhân", "field": "date_of_birth", "ap_dung": "tat_ca"},
    {"nhom": "Cá nhân", "field": "gender", "ap_dung": "tat_ca"},
    {"nhom": "Cá nhân", "field": "cell_number", "ap_dung": "tat_ca"},
    {"nhom": "Cá nhân", "field": "custom_zalo", "ap_dung": "tat_ca"},
    {"nhom": "Cá nhân", "field": "permanent_address", "ap_dung": "tat_ca"},
    {"nhom": "Công việc", "field": "department", "ap_dung": "tat_ca"},
    {"nhom": "Công việc", "field": "date_of_joining", "ap_dung": "tat_ca"},
    {"nhom": "Công việc", "field": "vn_contract_type", "ap_dung": "tat_ca"},
    {"nhom": "Công việc", "field": "contract_end_date", "ap_dung": "tat_ca"},
    {"nhom": "Công việc", "field": "attendance_device_id", "ap_dung": "tru_khoan"},
    {"nhom": "Lương", "field": common.F_LUONG_CB, "ap_dung": "tru_khoan"},
    {"nhom": "BHXH", "field": "ns_dong_bhxh", "ap_dung": "tat_ca"},
    {"nhom": "BHXH", "field": common.F_BAO_HIEM, "ap_dung": "dong_bhxh"},
    {"nhom": "BHXH", "field": "ns_so_so_bhxh", "ap_dung": "dong_bhxh"},
    {"nhom": "BHXH", "field": "vn_insurance_salary", "ap_dung": "dong_bhxh"},
    {"nhom": "BHXH", "field": "ns_bhxh_tu_ngay", "ap_dung": "dong_bhxh"},
    {"nhom": "BHXH", "field": "ns_ly_do_khong_bhxh", "ap_dung": "khong_bhxh"},
    {"nhom": "Thuế", "field": "vn_tax_code", "ap_dung": "tat_ca"},
    {"nhom": "Thuế", "field": "vn_number_of_dependents", "ap_dung": "tat_ca"},
    {"nhom": "Ngân hàng", "field": "bank_ac_no", "ap_dung": "co_luong_tk"},
    {"nhom": "Ngân hàng", "field": "bank_name", "ap_dung": "co_luong_tk"},
    {"nhom": "Ngân hàng", "field": "custom_bankowner", "ap_dung": "co_luong_tk"},
]

# Giấy tờ chỉ đếm (không nhập qua lưới)
GIAY_TO_KIEM = [
    {"ma": "gt_cccd", "label": "Giấy tờ: CCCD", "loai": {"CCCD"}, "con_han": False},
    {"ma": "gt_hdld", "label": "Giấy tờ: HĐLĐ còn hạn", "loai": {"Hợp đồng lao động", "Phụ lục HĐLĐ"}, "con_han": True},
    {"ma": "gt_ksk", "label": "Giấy tờ: Khám SK còn hạn", "loai": {"Giấy khám sức khỏe"}, "con_han": True},
    {"ma": "gt_attp", "label": "Giấy tờ: ATTP còn hạn", "loai": {"Xác nhận kiến thức ATTP"}, "con_han": True},
]

LABEL_MAC_DINH = {
    "custom_cccd": "CCCD", "date_of_birth": "Ngày sinh", "gender": "Giới tính",
    "cell_number": "SĐT", "custom_zalo": "Zalo", "permanent_address": "Thường trú",
    "department": "Bộ phận", "date_of_joining": "Ngày vào", "vn_contract_type": "Loại HĐ",
    "contract_end_date": "Hết hạn HĐ", "attendance_device_id": "Mã máy",
    common.F_LUONG_CB: "Lương CB", "ns_dong_bhxh": "Đóng BHXH", common.F_BAO_HIEM: "Mức trừ BHXH",
    "ns_so_so_bhxh": "Số sổ BHXH", "vn_insurance_salary": "Lương đóng BH",
    "ns_bhxh_tu_ngay": "Đóng BHXH từ", "ns_ly_do_khong_bhxh": "Lý do không BHXH",
    "vn_tax_code": "MST", "vn_number_of_dependents": "Số người phụ thuộc",
    "bank_ac_no": "Số TK", "bank_name": "Ngân hàng", "custom_bankowner": "Chủ TK",
}


def _meta_cot():
    """Mô tả cột cho client: field, label, fieldtype, options, nhom, ap_dung."""
    meta = common.meta_emp()
    tc = common.tuy_chon()
    cols = []
    for c in TRUONG_TOI_THIEU:
        f = meta.get_field(c["field"])
        if not f:
            continue  # site chưa có field (chưa migrate / chưa cài erpnextvn) → bỏ cột
        ft = f.fieldtype
        opts = None
        if ft == "Select":
            opts = [o for o in str(f.options or "").split("\n") if o != ""]
        elif ft == "Link":
            opts = tc.get(c["field"]) or None
        cols.append({
            "field": c["field"], "label": LABEL_MAC_DINH.get(c["field"], f.label or c["field"]),
            "fieldtype": ft, "options": opts, "nhom": c["nhom"], "ap_dung": c["ap_dung"],
        })
    return cols


def _ap_dung(col, emp):
    a = col["ap_dung"]
    if a == "tat_ca":
        return True
    khoan = common.la_khoan(emp.get("department"))
    if a == "tru_khoan":
        return not khoan
    dong = bool(cint(emp.get("ns_dong_bhxh")))
    if a == "dong_bhxh":
        return dong
    if a == "khong_bhxh":
        return not dong
    if a == "co_luong_tk":
        return flt(emp.get(common.F_LUONG_TK)) > 0
    return True


def _co_gia_tri(col, v):
    if col["fieldtype"] == "Check":
        return True  # cờ 0/1 luôn "có" (0 là lựa chọn hợp lệ)
    if col["fieldtype"] in ("Currency", "Float", "Int"):
        return v not in (None, "") and flt(v) != 0
    return v not in (None, "")


def _giay_to_du(gt_rows, kiem, hom_nay):
    for r in gt_rows:
        if common.nfc(r.get("loai")) not in {common.nfc(x) for x in kiem["loai"]}:
            continue
        if not kiem["con_han"]:
            return True
        hh = r.get("ngay_het_han")
        if hh and getdate(hh) >= hom_nay:
            return True
    return False


def tinh_thieu(emp, gt_rows, cols=None, hom_nay=None):
    """Trả (thieu[], ap_dung_count, du_count, giay_to{ma: bool})."""
    cols = cols if cols is not None else _meta_cot()
    hom_nay = hom_nay or getdate(today())
    thieu, n_ap, n_du = [], 0, 0
    for c in cols:
        if not _ap_dung(c, emp):
            continue
        n_ap += 1
        if _co_gia_tri(c, emp.get(c["field"])):
            n_du += 1
        else:
            thieu.append(c["field"])
    gt = {}
    for k in GIAY_TO_KIEM:
        n_ap += 1
        ok = _giay_to_du(gt_rows or [], k, hom_nay)
        gt[k["ma"]] = ok
        if ok:
            n_du += 1
        else:
            thieu.append(k["ma"])
    return thieu, n_ap, n_du, gt


def pct(n_du, n_ap):
    return round(100.0 * n_du / n_ap) if n_ap else 100


def _nhan_vien_active(department=None):
    filters = common.loc_bo_phan({"status": "Active"}, department)
    return common.lay_nhan_vien(filters)


@frappe.whitelist()
def tong_quan(department=None):
    """Tiến độ hồ sơ: mỗi NV Active → thieu[], pct_du; tong_hop đếm thiếu theo field."""
    guard()
    cols = _meta_cot()
    emps = _nhan_vien_active(department)
    gt = common.lay_giay_to([e.name for e in emps])
    hom_nay = getdate(today())
    dong, dem = [], {}
    tong_pct, so_du = 0, 0
    for e in emps:
        thieu, n_ap, n_du, _g = tinh_thieu(e, gt.get(e.name, []), cols, hom_nay)
        p = pct(n_du, n_ap)
        tong_pct += p
        if not thieu:
            so_du += 1
        for f in thieu:
            dem[f] = dem.get(f, 0) + 1
        dong.append({
            "name": e.name, "employee_name": e.employee_name, "department": e.department,
            "thieu": thieu, "pct_du": p,
        })
    labels = {c["field"]: c["label"] for c in cols}
    labels.update({k["ma"]: k["label"] for k in GIAY_TO_KIEM})
    thu_tu = [c["field"] for c in cols] + [k["ma"] for k in GIAY_TO_KIEM]
    return {
        "dong": dong,
        "tong_hop": {
            "so_nv": len(emps),
            "so_du": so_du,
            "trung_binh_pct": round(tong_pct / len(emps)) if emps else 100,
            "thieu_theo_field": [{"field": f, "label": labels.get(f, f), "so_luong": dem.get(f, 0)} for f in thu_tu if dem.get(f)],
        },
        "cot": cols,
        "giay_to": GIAY_TO_KIEM and [{"ma": k["ma"], "label": k["label"]} for k in GIAY_TO_KIEM],
    }


@frappe.whitelist()
def luoi(department=None, chi_thieu=True):
    """Ma trận NV × field cho lưới sửa inline. Mỗi ô {gia_tri, thieu, ap_dung}."""
    guard()
    chi_thieu = cint(chi_thieu) if not isinstance(chi_thieu, bool) else chi_thieu
    cols = _meta_cot()
    emps = _nhan_vien_active(department)
    gt = common.lay_giay_to([e.name for e in emps])
    hom_nay = getdate(today())
    dong = []
    for e in emps:
        thieu, n_ap, n_du, g = tinh_thieu(e, gt.get(e.name, []), cols, hom_nay)
        if chi_thieu and not thieu:
            continue
        o = {}
        for c in cols:
            ap = _ap_dung(c, e)
            v = e.get(c["field"])
            o[c["field"]] = {"gia_tri": v, "ap_dung": ap, "thieu": bool(ap and c["field"] in thieu)}
        dong.append({
            "name": e.name, "employee_name": e.employee_name, "department": e.department,
            "o": o, "giay_to": g, "thieu": thieu, "pct_du": pct(n_du, n_ap),
        })
    return {"cot": cols, "giay_to": [{"ma": k["ma"], "label": k["label"]} for k in GIAY_TO_KIEM], "dong": dong}


@frappe.whitelist()
def luu_o(employee, field, gia_tri=None):
    """Lưu 1 ô (whitelist §4.1.1, validate như luu). Qua get_doc().save() để hook chạy."""
    guard()
    if field not in common.WHITELIST_SUA or field == "ns_giay_to":
        frappe.throw(_("Field {0} không được sửa từ lưới.").format(field))
    doc = frappe.get_doc("Employee", employee)
    common.ap_dung_data(doc, {field: gia_tri})
    doc.save(ignore_permissions=True)
    frappe.db.commit()
    doc = frappe.get_doc("Employee", employee)
    emp = frappe._dict({f: doc.get(f) for f in common.fields_co("Employee", common.FIELD_DANH_SACH)})
    emp["name"] = doc.name
    gt = [r.as_dict() for r in (doc.get("ns_giay_to") or [])]
    cols = _meta_cot()
    thieu, n_ap, n_du, g = tinh_thieu(emp, gt, cols)
    o = {}
    for c in cols:
        ap = _ap_dung(c, emp)
        o[c["field"]] = {"gia_tri": emp.get(c["field"]), "ap_dung": ap, "thieu": bool(ap and c["field"] in thieu)}
    return {"name": doc.name, "thieu": thieu, "pct_du": pct(n_du, n_ap), "o": o, "giay_to": g,
            "gia_tri": doc.get(field)}


# ─────────── Excel ───────────

def _excel_cols():
    cols = _meta_cot()
    return [{"field": "name", "label": "Mã NV", "fieldtype": "Data"},
            {"field": "employee_name", "label": "Họ tên", "fieldtype": "Data"},
            {"field": "department", "label": "Bộ phận", "fieldtype": "Link"}] + cols


def _gia_tri_excel(col, v):
    if v in (None, ""):
        return None
    ft = col["fieldtype"]
    if ft == "Date":
        try:
            return getdate(v).strftime("%d/%m/%Y")
        except Exception:
            return str(v)
    if ft == "Check":
        return 1 if cint(v) else 0
    if ft in ("Currency", "Float"):
        return flt(v)
    if ft == "Int":
        return cint(v)
    return v


@frappe.whitelist()
def mau_excel(department=None):
    """Tải mẫu Excel: 1 dòng/NV Active; hàng 1 label, hàng 2 fieldname (xám); ô thiếu tô vàng."""
    guard()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = _excel_cols()
    emps = _nhan_vien_active(department)
    gt = common.lay_giay_to([e.name for e in emps])
    hom_nay = getdate(today())
    wb = Workbook()
    ws = wb.active
    ws.title = "Bo sung"
    vang = PatternFill("solid", fgColor="FFF59D")
    xam = Font(color="9E9E9E", italic=True, size=9)
    dam = Font(bold=True)
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=c["label"])
        cell.font = dam
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row=2, column=j, value=c["field"]).font = xam
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(28, len(c["label"]) + 4))
    ws.freeze_panes = "D3"
    meta_cols = cols[3:]
    for i, e in enumerate(emps, start=3):
        thieu, _a, _d, _g = tinh_thieu(e, gt.get(e.name, []), meta_cols, hom_nay)
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=i, column=j, value=_gia_tri_excel(c, e.get(c["field"])))
            if c["field"] in thieu:
                cell.fill = vang
    # Data validation list cho Select / Link có options
    n = max(len(emps) + 2, 3)
    for j, c in enumerate(cols, start=1):
        opts = c.get("options")
        if opts and c["fieldtype"] in ("Select", "Link") and sum(len(o) + 1 for o in opts) < 250:
            dv = DataValidation(type="list", formula1='"' + ",".join(o.replace('"', "") for o in opts) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(j)}3:{get_column_letter(j)}{n}")
        if c["fieldtype"] == "Check":
            dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(j)}3:{get_column_letter(j)}{n}")
    ws2 = wb.create_sheet("Huong dan")
    for k, line in enumerate([
        "Hàng 1 = tên cột, hàng 2 = fieldname (đừng xoá, dùng để nhập lại không lệch).",
        "Khớp nhân viên theo cột A (Mã NV). Không đổi cột A.",
        "Ô để trống = KHÔNG đổi dữ liệu đang có. Ô vàng = đang thiếu.",
        "Ngày: dd/mm/yyyy hoặc yyyy-mm-dd. Số tiền: bỏ dấu chấm/phẩy. Check: 1/0/x/có/không.",
    ], start=1):
        ws2.cell(row=k, column=1, value=line)
    bio = io.BytesIO()
    wb.save(bio)
    ten_bp = (department or "tat-ca").replace("__none__", "chua-co-bo-phan")
    frappe.response["filename"] = f"ns_bo_sung_{frappe.scrub(ten_bp)}_{today()}.xlsx"
    frappe.response["filecontent"] = bio.getvalue()
    frappe.response["type"] = "download"


def _doc_file(file_url):
    name = frappe.db.get_value("File", {"file_url": file_url}, "name")
    if not name:
        frappe.throw(_("Không tìm thấy file đã tải lên."))
    f = frappe.get_doc("File", name)
    return f.get_content()


def _parse_ngay(v):
    if v in (None, ""):
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    raise ValueError(f"Ngày không hợp lệ: {s}")


def _parse_o(col, v):
    """Chuỗi Excel → giá trị field. Trả (value, None) hoặc (None, loi)."""
    ft = col["fieldtype"]
    try:
        if ft == "Date":
            return _parse_ngay(v), None
        if ft == "Check":
            s = str(v).strip().lower()
            if s in ("1", "x", "có", "co", "yes", "true", "1.0"):
                return 1, None
            if s in ("0", "không", "khong", "no", "false", "0.0"):
                return 0, None
            return None, f"Giá trị cờ không hợp lệ: {v}"
        if ft in ("Currency", "Float"):
            s = re.sub(r"[.,\s]", "", str(v)) if isinstance(v, str) else v
            return flt(s), None
        if ft == "Int":
            return cint(re.sub(r"[.,\s]", "", str(v)) if isinstance(v, str) else v), None
        if ft in ("Select", "Link") and col.get("options"):
            s = str(v).strip()
            for o in col["options"]:
                if common.nfc(o).lower() == common.nfc(s).lower():
                    return o, None
            return None, f"'{s}' không nằm trong danh sách {col['label']}"
        s = str(v).strip()
        if col["field"] == "custom_cccd":
            s = re.sub(r"\D", "", s.replace(".0", "") if s.endswith(".0") else s)
            if not common.CCCD_RE.match(s):
                return None, "CCCD phải đúng 12 chữ số"
        return s, None
    except Exception as e:  # noqa
        return None, str(e)


def _bang_bang(v):
    return "" if v in (None, "") else str(v)


def _doc_excel(file_url):
    """Đọc file đã upload → (cols đã map theo hàng 2 hoặc hàng 1, rows)."""
    from openpyxl import load_workbook

    content = _doc_file(file_url)
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        frappe.throw(_("File không có dữ liệu."))
    all_cols = _excel_cols()
    by_field = {c["field"]: c for c in all_cols}
    by_label = {common.nfc(c["label"]).lower(): c for c in all_cols}
    header1 = [str(x).strip() if x is not None else "" for x in rows[0]]
    header2 = [str(x).strip() if x is not None else "" for x in rows[1]]
    co_hang_fieldname = any(h in by_field for h in header2)
    col_map = []
    for j, h in enumerate(header2 if co_hang_fieldname else header1):
        c = by_field.get(h) or by_label.get(common.nfc(h).lower())
        col_map.append(c)
    data_rows = rows[2:] if co_hang_fieldname else rows[1:]
    return col_map, data_rows


@frappe.whitelist()
def xem_truoc_excel(file_url):
    """Đọc file đã upload, khớp NV theo Mã NV (cột A), trả diff từng NV."""
    guard()
    col_map, data_rows = _doc_excel(file_url)
    if not col_map or not col_map[0] or col_map[0]["field"] != "name":
        frappe.throw(_("Cột A phải là Mã NV (name). Dùng file mẫu tải từ app."))
    names = [str(r[0]).strip() for r in data_rows if r and r[0] not in (None, "")]
    emps = {e.name: e for e in common.lay_nhan_vien({"name": ("in", names)})} if names else {}
    out, so_dong_loi = [], 0
    for r in data_rows:
        if not r or r[0] in (None, ""):
            continue
        ma = str(r[0]).strip()
        e = emps.get(ma)
        if not e:
            out.append({"employee": ma, "employee_name": None, "thay_doi": [], "loi": f"Không có nhân viên mã {ma}"})
            so_dong_loi += 1
            continue
        thay_doi, co_loi = [], False
        for j, c in enumerate(col_map):
            if j == 0 or not c or c["field"] in ("name", "employee_name"):
                continue
            v = r[j] if j < len(r) else None
            if v in (None, ""):
                continue  # ô trống = không đổi
            moi, loi = _parse_o(c, v)
            cu = e.get(c["field"])
            if loi:
                co_loi = True
                thay_doi.append({"field": c["field"], "label": c["label"], "cu": cu, "moi": str(v), "loi": loi})
                continue
            if _bang_bang(_gia_tri_so_sanh(c, cu)) == _bang_bang(_gia_tri_so_sanh(c, moi)):
                continue
            thay_doi.append({"field": c["field"], "label": c["label"], "cu": cu, "moi": moi, "loi": None})
        if co_loi:
            so_dong_loi += 1
        if thay_doi:
            out.append({"employee": ma, "employee_name": e.employee_name, "thay_doi": thay_doi, "loi": None})
    return {"dong": out, "so_dong_loi": so_dong_loi, "so_nv_thay_doi": len([d for d in out if d["thay_doi"]])}


def _gia_tri_so_sanh(col, v):
    if v in (None, ""):
        return ""
    ft = col["fieldtype"]
    if ft == "Date":
        try:
            return str(getdate(v))
        except Exception:
            return str(v)
    if ft in ("Currency", "Float"):
        return str(flt(v))
    if ft in ("Int", "Check"):
        return str(cint(v))
    return common.nfc(str(v)).strip()


@frappe.whitelist()
def ap_dung_excel(file_url, chi_dong=None):
    """Áp thay đổi đã xem trước. chi_dong = list mã NV (tuỳ chọn). Mỗi NV 1 save + commit;
    lỗi 1 NV không chặn NV khác. Tối đa 200 NV."""
    guard()
    chi_dong = set(common.parse_list(chi_dong))
    xem = xem_truoc_excel(file_url)
    thanh_cong, loi = 0, []
    dong = [d for d in xem["dong"] if d["thay_doi"] and (not chi_dong or d["employee"] in chi_dong)]
    if len(dong) > 200:
        frappe.throw(_("Tối đa 200 nhân viên mỗi lần nhập."))
    for d in dong:
        if any(t["loi"] for t in d["thay_doi"]):
            loi.append({"employee": d["employee"], "msg": "; ".join(t["loi"] for t in d["thay_doi"] if t["loi"])})
            continue
        try:
            doc = frappe.get_doc("Employee", d["employee"])
            common.ap_dung_data(doc, {t["field"]: t["moi"] for t in d["thay_doi"]})
            doc.save(ignore_permissions=True)
            frappe.db.commit()
            thanh_cong += 1
        except Exception as e:  # noqa
            frappe.db.rollback()
            loi.append({"employee": d["employee"], "msg": str(e)})
    return {"thanh_cong": thanh_cong, "loi": loi}
